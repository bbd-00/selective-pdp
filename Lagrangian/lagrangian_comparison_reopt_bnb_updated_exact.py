# lagrangian_comparison_reopt_bnb_updated_exact.py
#
# Compares three methods across all instances:
#   1. Exact     — full binary MILP solved by Gurobi (optimal)
#   2. BnB_Fresh — Branch-and-Bound, u=0 at every node, 100 subgradient iters
#   3. Reopt     — Subgradient + duplicate resolution + re-optimise + regret insertion
#
# Output: lagrangian_bnb_reopt_comparison_updated_exact.xlsx
#   Sheet "Exact"     — per-instance exact MIP results
#   Sheet "BnB_Fresh" — per-instance BnB_Fresh results
#   Sheet "Reopt"     — per-instance Reopt results
#   Sheet "Comparison" — all three side-by-side per instance

import os
import sys
import math
import time

import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))   # repo root: solve_exact_mip, instance_loader
from instance_loader        import load_instance_from_txt
from lagrangian_bnb_fresh import solve_instance_bnb
from lagrangian_reopt     import solve_instance_reopt
from solve_exact_mip      import solve_exact_mip

# ------------------------------------------------------------------
# Excel helpers
# ------------------------------------------------------------------
def autosize_columns(ws):
    for col in range(1, ws.max_column + 1):
        max_len    = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


DETAIL_HEADER = [
    "instance_file", "nbr_req", "nbr_cdr", "nbr_odr",
    "UB_LP", "best_L", "best_primal", "best_gap (%)",
    "iters_done", "total_time (s)", "avg_subproblem_time (s)",
    "n_nodes", "n_pruned", "n_feasible", "max_depth",
    "status",
]

EXACT_HEADER = [
    "instance_file", "nbr_req", "nbr_cdr", "nbr_odr",
    "exact_obj", "mip_gap (%)", "mip_status", "total_time (s)",
]

COMPARISON_HEADER = [
    "instance_file", "nbr_req", "nbr_cdr", "nbr_odr",
    # objectives
    "Exact_obj",
    "BnB_Fresh_primal", "Reopt_primal",
    # gap vs exact
    "BnB_Fresh_gap_vs_exact (%)", "Reopt_gap_vs_exact (%)",
    # lagrangian bound
    "BnB_Fresh_bestL", "Reopt_bestL",
    # runtime
    "Exact_time (s)", "BnB_Fresh_time (s)", "Reopt_time (s)",
    # BnB tree stats
    "BnB_Fresh_nodes", "BnB_Fresh_pruned", "BnB_Fresh_feasible",
    # reopt-specific
    "Reopt_n_dup_resolved", "Reopt_n_inserted",
]


def _detail_row(res, nbr_req, nbr_cdr, nbr_odr, instance_file):
    return [
        instance_file, nbr_req, nbr_cdr, nbr_odr,
        res.get("UB_LP"),
        res.get("best_L"),
        res.get("best_primal"),
        res.get("best_gap"),
        res.get("iters_done"),
        res.get("total_time"),
        res.get("avg_subproblem_time"),
        res.get("n_nodes"),
        res.get("n_pruned"),
        res.get("n_feasible"),
        res.get("max_depth"),
        res.get("status_msg"),
    ]


def _gap_vs_exact(primal, exact_obj):
    if primal is None or exact_obj is None or abs(exact_obj) < 1e-9:
        return None
    return 100.0 * (exact_obj - primal) / abs(exact_obj)


def _err_result(method_name):
    return {"status_msg": f"ERROR in {method_name}", "total_time": 0.0}


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    out_xlsx   = os.path.join(
        script_dir, "lagrangian_bnb_reopt_comparison_updated_exact.xlsx"
    )

    wb          = Workbook()
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    def _style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

    ws_exact   = wb.active
    ws_exact.title = "Exact"
    ws_exact.append(EXACT_HEADER)
    _style_header(ws_exact)

    ws_fresh = wb.create_sheet("BnB_Fresh")
    ws_fresh.append(DETAIL_HEADER)
    _style_header(ws_fresh)

    ws_reopt = wb.create_sheet("Reopt")
    ws_reopt.append(DETAIL_HEADER)
    _style_header(ws_reopt)

    ws_comp    = wb.create_sheet("Comparison")
    ws_comp.append(COMPARISON_HEADER)
    _style_header(ws_comp)

    # ------------------------------------------------------------------
    # Instance list
    # ------------------------------------------------------------------
    instances = []
    nbr_req = 30
    while nbr_req <= 60:
        for nbr_cdr in range(0, 12, 2):
            nbr_odr = 10 - nbr_cdr
            instance_file = (
                f"InstanceswithDemand/Instance_REQ{nbr_req}_CDR{nbr_cdr}"
                f"_ODR{nbr_odr}_FS0_SS0.txt"
            )
            instances.append((nbr_req, nbr_cdr, nbr_odr, instance_file))
        nbr_req += 10

    total = len(instances)
    print(f"Will solve {total} instances with 3 methods: Exact, BnB_Fresh, Reopt")
    print(f"Writing Excel to: {out_xlsx}")
    wb.save(out_xlsx)

    for idx, (nbr_req, nbr_cdr, nbr_odr, instance_file) in enumerate(instances, start=1):
        instance_path = instance_file

        print(f"\n{'='*70}")
        print(f"Instance {idx}/{total}: REQ={nbr_req}, CDR={nbr_cdr}, ODR={nbr_odr}")
        print(f"File: {instance_path}")
        print(f"{'='*70}")

        skip_detail = [instance_file, nbr_req, nbr_cdr, nbr_odr] + [None] * 12 + ["FILE NOT FOUND"]
        skip_exact  = [instance_file, nbr_req, nbr_cdr, nbr_odr] + [None] * 4

        if not os.path.exists(instance_path):
            print("[SKIP] FILE NOT FOUND")
            ws_exact.append(skip_exact)
            ws_fresh.append(skip_detail)
            ws_reopt.append(skip_detail)
            ws_comp.append([instance_file, nbr_req, nbr_cdr, nbr_odr] + [None] * (len(COMPARISON_HEADER) - 4))
            for ws in (ws_exact, ws_fresh, ws_reopt, ws_comp):
                autosize_columns(ws)
            wb.save(out_xlsx)
            continue

        # ---- Method 1: Exact MIP ----
        print("\n--- Exact MIP (Gurobi) ---")
        try:
            res_exact = solve_exact_mip(instance_path, time_limit=3600, verbose=False)
        except Exception as exc:
            print(f"  [ERROR] {exc}")
            res_exact = _err_result("Exact")
        ws_exact.append([
            instance_file, nbr_req, nbr_cdr, nbr_odr,
            res_exact.get("obj"),
            res_exact.get("mip_gap"),
            res_exact.get("status"),
            res_exact.get("time"),
        ])
        t_exact = res_exact.get("time") or 0.0
        print(f"  obj={res_exact.get('obj')}  gap={res_exact.get('mip_gap')}%  "
              f"status={res_exact.get('status')}  time={t_exact:.1f}s")

        # ---- Method 2: BnB Fresh (100 iters every node) ----
        print("\n--- BnB_Fresh (u=0, 100 iters/node) ---")
        try:
            res_fresh = solve_instance_bnb(
                instance_path, max_iters_root=100, max_iters_node=100,
                bnb_time_limit=3600, verbose=True
            )
        except Exception as exc:
            print(f"  [ERROR] {exc}")
            res_fresh = _err_result("BnB_Fresh")
        ws_fresh.append(_detail_row(res_fresh, nbr_req, nbr_cdr, nbr_odr, instance_file))
        print(f"  primal={res_fresh.get('best_primal')}  nodes={res_fresh.get('n_nodes')}  "
              f"time={res_fresh.get('total_time', 0.0):.1f}s")

        # ---- Method 3: Reopt ----
        print("\n--- Reopt ---")
        try:
            res_reopt = solve_instance_reopt(
                instance_path, max_iters=100, verbose=True
            )
        except Exception as exc:
            print(f"  [ERROR] {exc}")
            res_reopt = _err_result("Reopt")
        ws_reopt.append(_detail_row(res_reopt, nbr_req, nbr_cdr, nbr_odr, instance_file))
        print(f"  primal={res_reopt.get('best_primal')}  "
              f"n_dup_resolved={res_reopt.get('n_duplicates_resolved')}  "
              f"n_inserted={res_reopt.get('n_inserted')}  "
              f"time={res_reopt.get('total_time', 0.0):.1f}s")

        # ---- Comparison row ----
        exact_obj = res_exact.get("obj")
        ws_comp.append([
            instance_file, nbr_req, nbr_cdr, nbr_odr,
            exact_obj,
            res_fresh.get("best_primal"),
            res_reopt.get("best_primal"),
            _gap_vs_exact(res_fresh.get("best_primal"), exact_obj),
            _gap_vs_exact(res_reopt.get("best_primal"), exact_obj),
            res_fresh.get("best_L"),
            res_reopt.get("best_L"),
            res_exact.get("time"),
            res_fresh.get("total_time"),
            res_reopt.get("total_time"),
            res_fresh.get("n_nodes"),
            res_fresh.get("n_pruned"),
            res_fresh.get("n_feasible"),
            res_reopt.get("n_duplicates_resolved"),
            res_reopt.get("n_inserted"),
        ])

        for ws in (ws_exact, ws_fresh, ws_reopt, ws_comp):
            autosize_columns(ws)
        wb.save(out_xlsx)
        print(f"[SAVED] {out_xlsx}")

    print(f"\n{'='*70}")
    print(f"All done! Results saved to: {out_xlsx}")
    print(f"{'='*70}")


if __name__ == "__main__":
    main()
