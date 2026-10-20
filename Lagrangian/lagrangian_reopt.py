# lagrangian_reopt.py
#
# Lagrangian relaxation + primal heuristic (no MIP repair).
# After the subgradient converges, a three-step heuristic builds a primal:
#
#   1. Duplicate resolution  — shortcut-saving auction: among claimants of a
#      duplicate request, keep the courier with the lowest shortcut-saving
#      (cheapest to keep) and shortcut the others.
#
#   2. Re-optimisation       — for each courier, run GraphSearch with u=0
#      restricted to their assigned requests.  Gives the optimal route
#      for that assignment without interference from multipliers.
#
#   3. Insertion of unserved — scan unserved requests in regret order
#      (highest gap between best and second-best courier delta first).
#      Greedily insert one at a time; stop when no profitable insertion exists.

import os
import math
import time
import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root: instance_loader
from instance_loader import load_instance_from_txt
from graph_search import GraphSearchSolver


def solve_instance_reopt(instance_path: str,
                         max_iters: int = 100,
                         sub_time_limit: int = 60,
                         lp_time_limit: int = 300,
                         lambda0: float = 2.0,
                         max_noimpv: int = 7,
                         lambda_min: float = 1e-4,
                         verbose: bool = True):
    """
    Lagrangian subgradient (GraphSearch subproblems) followed by a three-step
    primal heuristic: duplicate resolution → re-optimisation → unserved insertion.

    Return dict keys (compatible with other Lagrangian files):
        instance, solver, UB_LP, best_L, best_primal, best_gap, lag_gap,
        iters_done, total_time, avg_subproblem_time,
        n_duplicates_resolved, n_inserted, status_msg.
    """
    t0 = time.time()

    if verbose:
        print(f"\n=== Loading instance: {instance_path} ===")

    instance = load_instance_from_txt(instance_path)

    b = int(instance["b"])
    num_comm_drivers = len(instance["committed"])
    num_occ_drivers  = len(instance["occasional"])
    num_requests     = len(instance["requests"])

    P  = list(range(1, num_requests + 1))
    D  = [i + num_requests for i in P]
    KC = list(range(0, num_comm_drivers))
    KO = list(range(num_comm_drivers, num_comm_drivers + num_occ_drivers))
    n  = len(P)
    V  = [0] + P + D + [2 * n + 1]
    K  = KC + KO

    A1 = [(0, j) for j in P]
    A2 = [(i, j) for i in P + D for j in P + D if i != j]
    A3 = [(i, 2 * n + 1) for i in D] + [(0, 2 * n + 1)]

    coords = {}
    for req in instance["requests"]:
        pid = req["id"]
        coords[pid]                = req["pickup"]
        coords[pid + num_requests] = req["dest"]

    ok = ([d["location"] for d in instance["committed"]] +
          [d["location"] for d in instance["occasional"]])

    dij_A2 = {(i, j): math.dist(coords[i], coords[j]) for (i, j) in A2}
    dij_A3 = {(i, j): 0 for (i, j) in A3}
    dij    = {**dij_A2, **dij_A3}
    dijk   = {(i, j, k): math.dist(ok[k], coords[j]) for (i, j) in A1 for k in K}
    tij    = dij.copy()
    tijk   = dijk.copy()

    si = {0: 0, 2 * num_requests + 1: 0}
    ei, li, dem = {}, {}, {}
    for req in instance["requests"]:
        pid = req["id"]
        ei[pid], li[pid]                           = req["e_pickup"],  req["l_pickup"]
        ei[pid + num_requests], li[pid + num_requests] = req["e_dest"],    req["l_dest"]
        si[pid]                = req["delay_p"]
        si[pid + num_requests] = req["delay_d"]
        dem[pid], dem[pid + num_requests] = req["l_dem"], -req["l_dem"]
    dem[0] = 0
    dem[2 * n + 1] = 0

    Ek = ({d["id"]: d["e"] for d in instance["committed"]} |
          {d["id"]: d["e"] for d in instance["occasional"]})
    Lk = ({d["id"]: d["l"] for d in instance["committed"]} |
          {d["id"]: d["l"] for d in instance["occasional"]})
    Qk = ({d["id"]: d["cap"] for d in instance["committed"]} |
          {d["id"]: d["cap"] for d in instance["occasional"]})

    p      = {req["id"]: req["rev"] for req in instance["requests"]}   # per-request revenue
    wC     = instance["c_com_h"] / b
    cO     = instance["c_occ_o"]
    deltaC = instance["c_com_d"]
    deltaO = instance["c_occ_d"]

    z_UB_base = {
        (pickup, driver): (
            1 if (li[pickup] - si[pickup] >= max(dijk[0, pickup, driver] + Ek[driver], ei[pickup] - si[pickup])
                  and Lk[driver] >= ei[pickup + num_requests])
            else 0
        )
        for driver in K for pickup in P
    }

    x_UB = {
        (i, j): (0 if (ei[i] + dij[i, j] + si[j] > li[j]) or (i in D and j == i - num_requests) else 1)
        for (i, j) in A2
    }
    UBxIPC1 = {
        (n2, d2): 0 if (d2 in D and d2 - num_requests != n2 and x_UB.get((d2 - num_requests, n2), 0) == 0) else 1
        for (n2, d2) in A2
    }
    UBxIPC2 = {
        (p2, n2): 0 if (p2 in P and p2 + num_requests != n2 and x_UB.get((n2, p2 + num_requests), 0) == 0) else 1
        for (p2, n2) in A2
    }
    x_ind = {(i, j): x_UB[i, j] * UBxIPC1[i, j] * UBxIPC2[i, j] for (i, j) in A2}

    A2 = [(i, j) for (i, j) in A2 if x_ind[i, j] == 1]
    A  = A1 + A2 + A3
    PK = [(pp, kk) for pp in P for kk in K if z_UB_base[pp, kk] == 1]
    Lk_org = Lk.copy()

    if verbose:
        print(f"Requests={num_requests}, CDR={num_comm_drivers}, ODR={num_occ_drivers}, Vehicles={len(K)}")
        print("Solving LP relaxation ...")

    # ------------------------------------------------------------------
    # LP Relaxation
    # ------------------------------------------------------------------
    def build_lp_model(timeLimit=300):
        m = gp.Model("LP_relax")
        m.Params.OutputFlag = 0
        if timeLimit:
            m.Params.TimeLimit = timeLimit

        x    = m.addVars(A, K, vtype=GRB.CONTINUOUS, name="x")
        Lvar = m.addVars(A2, vtype=GRB.CONTINUOUS, name="L")
        y    = m.addVars(P + D, K, vtype=GRB.CONTINUOUS, name="y")
        T    = m.addVars(P + D + [0, 2 * n + 1], K, vtype=GRB.CONTINUOUS, lb=0, name="T")

        obj_expr = gp.quicksum(p[i] * gp.quicksum(y[i, k] for k in K) for i in P)
        for k in KC:
            obj_expr -= wC * (Lk_org[k] - Ek[k]) + deltaC * (
                gp.quicksum(dijk[0, j, k] * x[0, j, k] for (_, j) in A1) +
                gp.quicksum(dij[i, j]     * x[i, j, k] for (i, j) in A2)
            )
        for k in KO:
            obj_expr -= cO * gp.quicksum(y[i, k] for i in P) + deltaO * gp.quicksum(
                dij[i, j] * x[i, j, k] for (i, j) in A2
            )
        m.setObjective(obj_expr, GRB.MAXIMIZE)

        for i in P:
            m.addConstr(gp.quicksum(x[i, j, k] for k in K for j in V if (i, j) in A) <= 1)
        for i in P + D:
            for k in K:
                m.addConstr(
                    gp.quicksum(x[j, i, k] for j in V if (j, i) in A) -
                    gp.quicksum(x[i, j, k] for j in V if (i, j) in A) == 0
                )
        for i in P:
            for k in K:
                m.addConstr(
                    gp.quicksum(x[i, j, k] for j in V if (i, j) in A) -
                    gp.quicksum(x[j, i + num_requests, k] for j in V if (j, i + num_requests) in A) == 0
                )
        for k in K:
            m.addConstr(gp.quicksum(x[0, j, k] for j in V if (0, j) in A) == 1)
            m.addConstr(gp.quicksum(x[j, 2 * n + 1, k] for j in V if (j, 2 * n + 1) in A) == 1)
        for i in P + D:
            for k in K:
                m.addConstr(y[i, k] == gp.quicksum(x[i, l, k] for l in V if (i, l) in A))
        for i in P:
            for k in K:
                if z_UB_base[i, k] == 0:
                    m.addConstr(y[i, k] == 0)

        svMTZ = {(i, j): min(-si[i] - dij[j, i], ei[j] - ei[i]) for (i, j) in A2}
        for (i, j) in A2:
            for k in K:
                if (j, i) in A2:
                    m.addConstr(
                        T[i,k] - T[j,k] + (tij[i,j]+si[j])*x[i,j,k] + svMTZ[i,j]*x[j,i,k]
                        <= li[i]*y[i,k] - ei[j]*y[j,k] - (li[i]-ei[j])*(x[i,j,k]+x[j,i,k])
                    )
                else:
                    m.addConstr(
                        T[i,k] - T[j,k] + (tij[i,j]+si[j])*x[i,j,k]
                        <= li[i]*y[i,k] - ei[j]*y[j,k] - (li[i]-ei[j])*x[i,j,k]
                    )
        for (i, j) in A3:
            if i != 0:
                for k in K:
                    m.addConstr(
                        T[i,k] - T[j,k] + (tij[i,j]+si[j])*x[i,j,k]
                        <= li[i]*y[i,k]
                        - Ek[k]*gp.quicksum(x[0,ii,k] for ii in P if (ii,k) in PK)
                        - (li[i]-Ek[k])*x[i,j,k]
                    )
        for (i, j) in A1:
            for k in K:
                m.addConstr(
                    T[0,k] - T[j,k] + (tijk[0,j,k]+si[j])*x[0,j,k]
                    <= Lk[k]*gp.quicksum(x[0,ii,k] for ii in P if (ii,k) in PK)
                    - ei[j]*y[j,k] - (Lk[k]-ei[j])*x[0,j,k]
                )
        for i in P + D:
            for k in K:
                m.addConstr(ei[i]*y[i,k] <= T[i,k])
                m.addConstr(T[i,k] <= li[i]*y[i,k])
        for i in P:
            for k in K:
                m.addConstr(T[i,k] + (tij[i,i+num_requests]+si[i+num_requests])*y[i,k]
                            <= T[i+num_requests,k])
        for k in K:
            m.addConstr(Ek[k]*gp.quicksum(x[0,i,k] for i in P if (i,k) in PK) <= T[0,k])
            m.addConstr(T[2*n+1,k] <= Lk[k]*gp.quicksum(x[0,i,k] for i in P if (i,k) in PK))
        # Load flow with one load variable per arc, shared by all couriers (each arc is used by at
        # most one courier, so the sums over k pick the capacity of whichever courier uses it).
        for i in P + D:
            m.addConstr(
                gp.quicksum(Lvar[i, j] for j in P + D if (i, j) in A2) -
                gp.quicksum(Lvar[j, i] for j in P + D if (j, i) in A2)
                == dem[i] * gp.quicksum(x[i, j, k] for k in K for j in V if (i, j) in A)
            )
        for (i, j) in A2:
            m.addConstr(Lvar[i, j] >= gp.quicksum(max(dem[i], 0) * x[i, j, k] for k in K))
            m.addConstr(Lvar[i, j] <= gp.quicksum((Qk[k] + min(dem[i], 0)) * x[i, j, k] for k in K))
        return m

    m_lp = build_lp_model(lp_time_limit)
    m_lp.optimize()

    if m_lp.status not in (GRB.OPTIMAL, GRB.TIME_LIMIT):
        if verbose:
            print(f"[LP] FAILED status={m_lp.status}")
        return {
            "instance": instance_path, "solver": "Reopt_GraphSearch",
            "UB_LP": None, "best_L": None, "best_primal": None,
            "best_gap": None, "lag_gap": None, "iters_done": 0,
            "total_time": time.time() - t0, "avg_subproblem_time": None,
            "n_duplicates_resolved": 0, "n_inserted": 0,
            "status_msg": f"LP failed status={m_lp.status}"
        }

    UB_LP = float(m_lp.objVal)
    if verbose:
        print(f"[LP] UB_LP = {UB_LP:.6f}")

    # ------------------------------------------------------------------
    # Lagrangian subgradient — also tracks per-courier routes
    # ------------------------------------------------------------------
    u            = {i: 0.0 for i in P}
    lambda_r     = float(lambda0)
    no_improve   = 0
    best_L       = float("inf")
    it_done      = 0
    served_by    = {i: {k: 0.0 for k in K} for i in P}
    routes       = {k: [] for k in K}
    sub_times    = []

    if verbose:
        print(f"Starting subgradient ...")
        print("Iter | L(u)           | best_L         | ||g||^2       | lambda   | noimp")

    for it in range(1, max_iters + 1):
        it_done         = it
        sum_sub_obj     = 0.0
        served_by       = {i: {k: 0.0 for k in K} for i in P}
        served_sum_by_i = {i: 0.0 for i in P}
        iter_times      = []

        for k in K:
            ts  = time.time()
            gs  = GraphSearchSolver(instance, k, z_UB_base, x_UB, x_ind, u, verbose=False)
            res = gs.solve()
            for pickup in res["served_pickups"]:
                served_sum_by_i[pickup] += 1.0
                served_by[pickup][k]    = 1.0
            routes[k] = res.get("route") or []
            iter_times.append(time.time() - ts)
            sum_sub_obj += res["objVal"]

        sub_times.extend(iter_times)
        lag_val   = sum_sub_obj + sum(u.values())
        g         = {i: 1.0 - served_sum_by_i[i] for i in P}
        norm_g_sq = sum(g[i] ** 2 for i in P)

        improved = lag_val < best_L - 1e-9
        if improved:
            best_L     = lag_val
            no_improve = 0
        else:
            no_improve += 1

        if verbose:
            avg_st = sum(iter_times) / len(iter_times)
            print(f"{it:4d} | {lag_val:14.6f} | {best_L:14.6f} | {norm_g_sq:12.6f} | {lambda_r:7.4f} | {no_improve:5d}")

        if norm_g_sq <= 1e-12:
            if verbose:
                print("  --> Zero subgradient. Stopping.")
            break

        step = lambda_r * (lag_val - (-abs(UB_LP))) / norm_g_sq
        for i in P:
            u[i] = max(0.0, u[i] - step * g[i])

        if no_improve >= max_noimpv:
            lambda_r   *= 0.5
            no_improve  = 0
            if verbose:
                print(f"  --> Reducing lambda to {lambda_r:.6f}")
            if lambda_r < lambda_min:
                if verbose:
                    print("  --> lambda too small. Stopping.")
                break

    n_duplicates = sum(1 for i in P if sum(1 for k in K if served_by[i][k] > 0.5) > 1)
    n_unserved   = sum(1 for i in P if sum(1 for k in K if served_by[i][k] > 0.5) == 0)
    if verbose:
        print(f"  [End Lagrangian] duplicates={n_duplicates}  unserved={n_unserved}  n={len(P)}")

    # ------------------------------------------------------------------
    # Primal heuristic helpers
    # ------------------------------------------------------------------
    def _arc_cost(a, b, k):
        if b == 2 * n + 1:
            return 0.0
        if a == 0:
            return dijk.get((0, b, k), 0.0)
        return dij.get((a, b), 0.0)

    def _shortcut_saving(i, k, route, n_assigned_k):
        """Distance-cost saved by removing request i from courier k's route."""
        pickup, delivery = i, i + n
        seq = [0] + list(route) + [2 * n + 1]
        if pickup not in seq or delivery not in seq:
            return 0.0

        pos_p     = seq.index(pickup)
        dist_saved = (_arc_cost(seq[pos_p-1], pickup, k) +
                      _arc_cost(pickup, seq[pos_p+1], k) -
                      _arc_cost(seq[pos_p-1], seq[pos_p+1], k))
        seq_no_p  = seq[:pos_p] + seq[pos_p+1:]

        pos_d      = seq_no_p.index(delivery)
        dist_saved += (_arc_cost(seq_no_p[pos_d-1], delivery, k) +
                       _arc_cost(delivery, seq_no_p[pos_d+1], k) -
                       _arc_cost(seq_no_p[pos_d-1], seq_no_p[pos_d+1], k))

        if k in KC:
            return deltaC * dist_saved
        remaining_after = n_assigned_k - 1
        return deltaO * dist_saved + (cO if remaining_after == 0 else 0.0)

    def _gs_for(k, allowed_set):
        """Run GraphSearch for courier k restricted to allowed_set, u=0."""
        u_zero = {i: 0.0 for i in P}
        z      = {(i2, k2): (z_UB_base.get((i2, k2), 0)
                              if k2 == k and i2 in allowed_set else 0)
                  for i2 in P for k2 in K}
        gs  = GraphSearchSolver(instance, k, z, x_UB, x_ind, u_zero, verbose=False)
        res = gs.solve()
        return res["objVal"], res.get("route") or []

    # ------------------------------------------------------------------
    # Step 1: Duplicate resolution via shortcut-saving auction
    # ------------------------------------------------------------------
    if verbose:
        print("\n[Step 1] Resolving duplicates ...")

    assignment = {k: set() for k in K}
    n_resolved = 0

    for i in P:
        claimants = [k for k in K if served_by[i][k] > 0.5]
        if len(claimants) == 0:
            pass
        elif len(claimants) == 1:
            assignment[claimants[0]].add(i)
        else:
            n_resolved += 1
            n_assigned  = {k: sum(1 for j in P if served_by[j][k] > 0.5) for k in claimants}
            savings     = {k: _shortcut_saving(i, k, routes[k], n_assigned[k])
                           for k in claimants}
            keeper = min(claimants, key=lambda k: savings[k])
            assignment[keeper].add(i)

    if verbose:
        sizes = {k: len(assignment[k]) for k in K}
        print(f"  duplicates resolved={n_resolved}  assignment sizes={sizes}")

    # ------------------------------------------------------------------
    # Step 2: Re-optimise each courier's route with GraphSearch (u=0)
    # ------------------------------------------------------------------
    if verbose:
        print("\n[Step 2] Re-optimising routes ...")

    obj_vals = {}
    total_reopt = 0.0
    for k in K:
        obj_vals[k], routes[k] = _gs_for(k, assignment[k])
        total_reopt += obj_vals[k]

    if verbose:
        print(f"  primal after re-opt = {total_reopt:.6f}")

    # ------------------------------------------------------------------
    # Step 3: Insert unserved requests in regret order
    # ------------------------------------------------------------------
    if verbose:
        print("\n[Step 3] Inserting unserved requests ...")

    unserved  = [i for i in P if not any(i in assignment[k] for k in K)]
    n_inserted = 0

    inserted = True
    while inserted and unserved:
        inserted   = False
        candidates = []

        for i in unserved:
            best_delta  = -float("inf")
            best2_delta = -float("inf")
            best_k      = None
            best_obj    = None
            best_route  = None

            for k in K:
                if z_UB_base.get((i, k), 0) == 0:
                    continue
                new_obj, new_route = _gs_for(k, assignment[k] | {i})
                delta = new_obj - obj_vals[k]
                if delta > best_delta:
                    best2_delta = best_delta
                    best_delta  = delta
                    best_k      = k
                    best_obj    = new_obj
                    best_route  = new_route
                elif delta > best2_delta:
                    best2_delta = delta

            if best_k is not None and best_delta > 1e-6:
                regret = best_delta - max(best2_delta, 0.0)
                candidates.append((regret, best_delta, i, best_k, best_obj, best_route))

        if not candidates:
            break

        candidates.sort(reverse=True)
        _, delta, i_ins, k_ins, new_obj, new_route = candidates[0]

        assignment[k_ins].add(i_ins)
        obj_vals[k_ins] = new_obj
        routes[k_ins]   = new_route
        unserved.remove(i_ins)
        n_inserted += 1
        inserted    = True

        if verbose:
            print(f"  inserted req {i_ins} → courier {k_ins}  delta={delta:.4f}")

    best_primal = sum(obj_vals.values())

    if verbose:
        print(f"\n  primal after insertion = {best_primal:.6f}")
        print(f"  inserted={n_inserted}  still unserved={len(unserved)}")

    elapsed         = time.time() - t0
    avg_sub_time    = sum(sub_times) / len(sub_times) if sub_times else None
    best_gap        = (100.0 * (UB_LP - best_primal) / abs(best_primal)
                       if abs(best_primal) > 1e-9 else None)
    lag_gap         = (100.0 * (best_L - best_primal) / abs(best_primal)
                       if abs(best_primal) > 1e-9 else None)

    if verbose:
        print(f"\nDone. best_L={best_L:.6f}  UB_LP={UB_LP:.6f}  best_primal={best_primal:.6f}")
        print(f"lag_gap={lag_gap:.2f}%  UB_LP_gap={best_gap:.2f}%  time={elapsed:.1f}s")

    return {
        "instance":              instance_path,
        "solver":                "Reopt_GraphSearch",
        "UB_LP":                 UB_LP,
        "best_L":                float(best_L),
        "best_primal":           best_primal,
        "best_gap":              best_gap,
        "lag_gap":               lag_gap,
        "iters_done":            it_done,
        "lambda_final":          float(lambda_r),
        "total_time":            elapsed,
        "avg_subproblem_time":   avg_sub_time,
        "n_duplicates_resolved": n_resolved,
        "n_inserted":            n_inserted,
        "status_msg":            "OK",
    }


# ------------------------------------------------------------------
# Excel helpers
# ------------------------------------------------------------------
def autosize_columns(ws):
    for col in range(1, ws.max_column + 1):
        max_len    = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    out_xlsx   = os.path.join(script_dir, "lagrangian_reopt_results.xlsx")

    wb          = Workbook()
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    def _style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

    ws_sum = wb.active
    ws_sum.title = "Reopt_Summary"
    summary_header = [
        "instance_file", "nbr_req", "nbr_cdr", "nbr_odr",
        "UB_LP", "best_L", "best_primal",
        "best_gap_vs_LP (%)", "lag_gap (%)",
        "iters_done", "lambda_final",
        "total_time (s)", "avg_subproblem_time (s)",
        "n_duplicates_resolved", "n_inserted", "status",
    ]
    ws_sum.append(summary_header)
    _style_header(ws_sum)

    instances = []
    nbr_req = 10
    while nbr_req <= 30:
        for nbr_cdr in range(0, 12, 2):
            nbr_odr       = 10 - nbr_cdr
            instance_file = (f"InstanceswithDemand/"
                             f"Instance_REQ{nbr_req}_CDR{nbr_cdr}_ODR{nbr_odr}_FS0_SS0.txt")
            instances.append((nbr_req, nbr_cdr, nbr_odr, instance_file))
        nbr_req += 10

    print(f"Will solve {len(instances)} instances with Lagrangian + Reopt heuristic")
    print(f"Writing Excel to: {out_xlsx}")
    wb.save(out_xlsx)

    total = len(instances)

    for idx, (nbr_req, nbr_cdr, nbr_odr, instance_file) in enumerate(instances, start=1):
        print(f"\n{'='*70}")
        print(f"Instance {idx}/{total}: REQ={nbr_req}, CDR={nbr_cdr}, ODR={nbr_odr}")
        print(f"File: {instance_file}")
        print(f"{'='*70}")

        if not os.path.exists(instance_file):
            print("[SKIP] FILE NOT FOUND")
            ws_sum.append([instance_file, nbr_req, nbr_cdr, nbr_odr] + [None] * 12)
            autosize_columns(ws_sum)
            wb.save(out_xlsx)
            continue

        res = solve_instance_reopt(instance_file, verbose=True)

        ws_sum.append([
            instance_file, nbr_req, nbr_cdr, nbr_odr,
            res.get("UB_LP"),
            res.get("best_L"),
            res.get("best_primal"),
            res.get("best_gap"),
            res.get("lag_gap"),
            res.get("iters_done"),
            res.get("lambda_final"),
            res.get("total_time"),
            res.get("avg_subproblem_time"),
            res.get("n_duplicates_resolved"),
            res.get("n_inserted"),
            res.get("status_msg"),
        ])

        autosize_columns(ws_sum)
        wb.save(out_xlsx)
        print(f"[SAVED] {out_xlsx}")


if __name__ == "__main__":
    main()
