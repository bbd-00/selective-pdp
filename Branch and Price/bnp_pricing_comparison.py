# bnp_pricing_comparison.py
#
# Compares Branch-and-Price pricing solvers (and optionally the exact MIP) on every
# instance file in INSTANCE_DIR or its subfolders whose name contains _REQ{n}_CDR{c}_ODR{o}
# with n ≤ MAX_REQ, ordered by subfolder, then REQ, CDR and file name.  The instance_file
# column shows each path relative to INSTANCE_DIR.  Methods are defined in the METHODS
# registry below — sheets, headers, and comparison columns are generated from it.
#
# Every method (and the exact MIP) is run once per time-window tightening mode in
# TIGHTEN_MODES; both solvers then use the same preprocessing (preprocessing.py), whose time is
# reported separately (prep_time) and is not part of any solve time.
#
# Output: bnp_pricing_comparison_{dataset}_{cgall|cgfirst}_{modes}_{twmodes}.xlsx, where
# {dataset} is the last folder of INSTANCE_DIR, {modes} lists the MIP-heuristic settings run
# (e.g. "mip_nomip") and {twmodes} the tightening settings (e.g. "notw_tw").
#   Sheet "Exact"        — exact MIP results, one row per instance and tightening mode
#                          (only if RUN_EXACT)
#   Sheet "<MethodName>" — one detail sheet per configured method
#   Sheet "Comparison"   — all methods side-by-side per instance, plus pairwise columns:
#                          every method vs the first (reference) method, every method vs its
#                          own "baselines" (e.g. BnP_GS_nomip_tw vs BnP_GS_mip_tw and vs
#                          BnP_GS_nomip_notw), and the tightened exact MIP vs the untightened one
#   Sheet "Solution"     — statistics of each method's final solution (see solution_stats.py),
#                          one row per instance and method

import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill

_HERE      = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _REPO_ROOT)   # solve_exact_mip, instance_loader
sys.path.insert(0, _HERE)        # branch_and_price and its pricing solvers

from branch_and_price import solve_instance_bnp
from instance_loader import load_instance_from_txt
from solution_stats import STAT_KEYS, solution_statistics
from solve_exact_mip import solve_exact_mip


# ── Tunable parameters ────────────────────────────────────────────────────────
INSTANCE_DIR      = "InstanceswithDemand"  # folder relative to the repo root, e.g.
                                           # "UpdatedInstances/converted_sun_120_rev"
MAX_REQ           = 100              # largest REQ size to include
RUN_EXACT         = False            # False: skip the exact MIP (two-way pricing comparison)
EXACT_TIME_LIMIT  = 3600            # seconds for the exact MIP
TIME_LIMIT        = 3600            # seconds per B&P run
MIP_TIME_PER_NODE = 20.0            # seconds for the RMP MIP at each node
MIP_HEURISTIC_MODES = [True]  # per-node MIP heuristic settings to run, each as its own
                                     # methods: [True] / [False] / [True, False] (side by side)
ADD_ALL_COLUMNS   = True            # False: first positive-rc column per round; True: all
TIGHTEN_MODES     = [False, True]   # time-window tightening settings to run, each as its own
                                    # methods: [False] / [True] / [False, True] (side by side)
# ─────────────────────────────────────────────────────────────────────────────

_BNP_KWARGS = {
    "time_limit":        TIME_LIMIT,
    "mip_time_per_node": MIP_TIME_PER_NODE,
    "add_all_columns":   ADD_ALL_COLUMNS,
    "verbose":           False,
}

_BNP_EXTRA_COLS = [
    ("best_ub",            "best_ub"),
    ("root_gap (%)",       "root_gap"),
    ("final_gap (%)",      "best_gap"),
    ("gap_improvement (%)", "gap_improvement"),
    ("nodes",              "n_nodes"),
    ("pruned",             "n_pruned"),
    ("lp_infeasible",      "n_lp_infeasible"),
    ("mip_improvements",   "n_mip_improvements"),
    ("cols_root",          "n_cols_root"),
    ("cols_total",         "n_cols_total"),
    ("cg_iters_root",      "n_cg_iters_root"),
    ("cg_iters_total",     "n_cg_iters_total"),
    ("cg_time (s)",        "cg_time"),
    ("pricing_time (s)",   "pricing_time"),
    ("pricing_calls",      "n_pricing_calls"),
    ("bound_checks",       "n_bound_checks"),
    ("bound_pruned",       "n_bound_pruned"),
    ("requests_eligible",  "n_requests_eligible"),
    ("requests_excluded",  "n_requests_excluded"),
    ("mip_time (s)",       "mip_time"),
]

# Preprocessing statistics (preprocessing.preprocess_instance), for B&P and the exact MIP
_PREP_COLS = [
    ("prep_time (s)",          "prep_time"),
    ("prep_rounds",            "prep_rounds"),
    ("tw_width_reduction (%)", "tw_width_reduction_servable_pct"),
    ("unservable_requests",    "unservable_requests"),
    ("zUB_pairs",              "zUB_pairs"),
    ("arcs_kept",              "arcs_kept"),
]


def _flatten_prep(res):
    """Copy the preprocessing statistics into the result dict so they can be read by key."""
    for key, value in (res.get("prep_stats") or {}).items():
        res.setdefault(key, value)
    return res


# ── Method registry ───────────────────────────────────────────────────────────
# The first entry is the reference for the "vs reference" comparison columns.
#
# Required keys per entry:
#   name             — label used for the sheet name and all column prefixes
#   func             — callable(instance_path, **kwargs) → dict
#   kwargs           — passed verbatim to func on every call
#   key_primal       — result-dict key for the best primal objective value
#   key_bound        — result-dict key for the bound compared across methods (or None)
#   key_time         — result-dict key for total wall-clock time
#   key_time_to_best — result-dict key for the time at which the final best solution was found
#   key_status       — result-dict key for the status / message string
#   key_pricing_time — result-dict key for pricing time (or None)
#   key_routes       — result-dict key for the solution's routes {courier: [nodes]} (or None),
#                      used for the "Solution" statistics sheet
#   extra_cols       — list of (column_label, result_dict_key) for method-specific stats
# Optional:
#   baselines        — names of other methods to compare this one against directly

_PRICING_SOLVERS = [("GS", "graphsearch"), ("Label", "labeling")]


def _mode_tag(use_mip):
    return "mip" if use_mip else "nomip"


def _tw_tag(tighten):
    return "tw" if tighten else "notw"


def _bnp_name(solver_tag, use_mip, tighten):
    return f"BnP_{solver_tag}_{_mode_tag(use_mip)}_{_tw_tag(tighten)}"


def _bnp_method(solver_tag, pricing_algo, use_mip, tighten):
    method = {
        "name":             _bnp_name(solver_tag, use_mip, tighten),
        "func":             solve_instance_bnp,
        "kwargs":           {**_BNP_KWARGS, "pricing_algo": pricing_algo,
                             "use_mip_heuristic": use_mip, "tighten": tighten},
        "key_primal":       "best_primal",
        "key_bound":        "lp_root_bound",
        "key_time":         "total_time",
        "key_time_to_best": "time_to_best",
        "key_status":       "status_msg",
        "key_pricing_time": "pricing_time",
        "key_routes":       "routes",
        "extra_cols":       _BNP_EXTRA_COLS + _PREP_COLS,
        "baselines":        [],
    }
    # Without the heuristic, compare against the same pricing solver with it.
    if not use_mip and True in MIP_HEURISTIC_MODES:
        method["baselines"].append(_bnp_name(solver_tag, True, tighten))
    # With tightening, compare against the same method without it.
    if tighten and False in TIGHTEN_MODES:
        method["baselines"].append(_bnp_name(solver_tag, use_mip, False))
    return method


METHODS = [_bnp_method(tag, algo, use_mip, tighten)
           for tag, algo in _PRICING_SOLVERS
           for use_mip in MIP_HEURISTIC_MODES
           for tighten in TIGHTEN_MODES]

# Exact MIP, once per tightening mode; the first one is the reference for gap_vs_exact
EXACT_RUNS = [(f"Exact_{_tw_tag(tighten)}", tighten) for tighten in TIGHTEN_MODES]
# ─────────────────────────────────────────────────────────────────────────────


# ── Excel helpers ─────────────────────────────────────────────────────────────
def _autosize(ws):
    for col in range(1, ws.max_column + 1):
        max_len    = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def _safe(v, decimals=4):
    if v is None:
        return None
    try:
        return round(float(v), decimals)
    except (TypeError, ValueError):
        return v


def _gap_vs_exact(primal, exact_obj):
    if primal is None or exact_obj is None or abs(exact_obj) < 1e-9:
        return None
    return round(100.0 * (exact_obj - primal) / abs(exact_obj), 4)


def _diff(a, b):
    if a is None or b is None:
        return None
    return round(a - b, 6)


def _ratio(ref, other):
    if ref is None or other is None or other < 1e-9:
        return None
    return round(ref / other, 3)


# ── Header builders ───────────────────────────────────────────────────────────
_INSTANCE_COLS = ["instance_file", "nbr_req", "nbr_cdr", "nbr_odr"]

EXACT_HEADER = _INSTANCE_COLS + [
    "tighten", "Exact_obj", "Exact_lp_relax_ub", "Exact_mip_gap (%)", "Exact_time (s)",
    "Exact_time_to_best (s)", "Exact_status",
] + [label for label, _ in _PREP_COLS]


def _exact_row(instance_cols, tighten, res):
    return list(instance_cols) + [
        tighten,
        _safe(res.get("obj")),
        _safe(res.get("lp_relax_ub")),
        _safe(res.get("mip_gap")),
        _safe(res.get("time"), 2),
        _safe(res.get("time_to_best"), 2),
        res.get("status"),
    ] + [_safe(res.get(rkey)) for _, rkey in _PREP_COLS]

_DETAIL_BASE = _INSTANCE_COLS + [
    "best_primal", "gap_vs_exact (%)", "bound", "total_time (s)", "time_to_best (s)", "status",
]


SOLUTION_HEADER = _INSTANCE_COLS + ["method", "status", "reported_profit"] + STAT_KEYS


def _solution_row(instance, instance_cols, method_name, status, reported_profit, routes):
    """One row of the "Solution" sheet: statistics of a method's final solution."""
    row = list(instance_cols) + [method_name, status, _safe(reported_profit)]
    if reported_profit is None or routes is None:
        return row + [None] * len(STAT_KEYS)
    stats = solution_statistics(instance, routes, reported_profit)
    return row + [_safe(stats[key]) for key in STAT_KEYS]


def _detail_header(method):
    return _DETAIL_BASE + [label for label, _ in method["extra_cols"]]


def _comparison_pairs(methods):
    """(method, baseline) pairs for the pairwise columns: every method vs the first one,
    then every method vs each of its own declared baselines, without duplicates."""
    pairs = [(m, methods[0]) for m in methods[1:]]
    by_name = {m["name"]: m for m in methods}
    for m in methods:
        for base_name in m.get("baselines", []):
            base = by_name.get(base_name)
            if base is not None and (m, base) not in pairs:
                pairs.append((m, base))
    return pairs


def _exact_names(run_exact):
    return [name for name, _ in EXACT_RUNS] if run_exact else []


def _comparison_header(methods, run_exact):
    ex = _exact_names(run_exact)
    cols = list(_INSTANCE_COLS)
    # ── objectives / primals ──
    cols += [f"{e}_obj" for e in ex]
    cols += [f"{m['name']}_primal" for m in methods]
    # ── gap vs exact (the first exact run is the reference) ──
    cols += [f"{e}_mip_gap (%)" for e in ex]
    if ex:
        cols += [f"{m['name']}_gap_vs_exact (%)" for m in methods]
    # ── bounds ──
    cols += [f"{e}_lp_relax_ub" for e in ex]
    cols += [f"{m['name']}_bound" for m in methods]
    # ── wall-clock times ──
    cols += [f"{e}_time (s)" for e in ex]
    cols += [f"{m['name']}_time (s)" for m in methods]
    # ── time until the final best solution was found ──
    cols += [f"{e}_time_to_best (s)" for e in ex]
    cols += [f"{m['name']}_time_to_best (s)" for m in methods]
    # ── statuses ──
    cols += [f"{e}_status" for e in ex]
    cols += [f"{m['name']}_status" for m in methods]
    # ── pairwise: each exact run vs the first one (tightening on vs off) ──
    for e in ex[1:]:
        cols += [f"{e}_obj_diff_vs_{ex[0]}",
                 f"{e}_lp_diff_vs_{ex[0]}",
                 f"{e}_time_speedup_vs_{ex[0]}"]
    # ── pairwise: method vs reference, method vs its own baselines ──
    for m, base in _comparison_pairs(methods):
        a, b = m["name"], base["name"]
        cols += [f"{a}_bound_diff_vs_{b}",
                 f"{a}_primal_diff_vs_{b}",
                 f"{a}_time_speedup_vs_{b}",
                 f"{a}_time_to_best_speedup_vs_{b}"]
        if m.get("key_pricing_time") and base.get("key_pricing_time"):
            cols.append(f"{a}_pricing_speedup_vs_{b}")
    # ── method-specific extras (grouped by method) ──
    for m in methods:
        cols += [f"{m['name']}_{label}" for label, _ in m["extra_cols"]]
    return cols


# ── Row builders ──────────────────────────────────────────────────────────────
def _detail_row(method, res, instance_cols, exact_obj):
    row = list(instance_cols) + [
        _safe(res.get(method["key_primal"])),
        _gap_vs_exact(res.get(method["key_primal"]), exact_obj),
        _safe(res.get(method["key_bound"])) if method["key_bound"] else None,
        _safe(res.get(method["key_time"]), 2),
        _safe(res.get(method.get("key_time_to_best")), 2) if method.get("key_time_to_best") else None,
        res.get(method["key_status"]),
    ]
    for _, rkey in method["extra_cols"]:
        row.append(_safe(res.get(rkey)))
    return row


def _comparison_row(methods, results, res_exact, instance_cols, run_exact):
    """res_exact: {exact run name: result dict} (empty if not run_exact)."""
    def get(m, key_field):
        key = m.get(key_field)
        return results[m["name"]].get(key) if key else None

    ex = _exact_names(run_exact)
    exr = [res_exact.get(e, {}) for e in ex]
    exact_obj = exr[0].get("obj") if exr else None
    row = list(instance_cols)
    # ── objectives / primals ──
    row += [_safe(r.get("obj")) for r in exr]
    row += [_safe(get(m, "key_primal")) for m in methods]
    # ── gap vs exact ──
    row += [_safe(r.get("mip_gap")) for r in exr]
    if ex:
        row += [_gap_vs_exact(get(m, "key_primal"), exact_obj) for m in methods]
    # ── bounds ──
    row += [_safe(r.get("lp_relax_ub")) for r in exr]
    row += [_safe(get(m, "key_bound")) for m in methods]
    # ── wall-clock times ──
    row += [_safe(r.get("time"), 2) for r in exr]
    row += [_safe(get(m, "key_time"), 2) for m in methods]
    # ── time until the final best solution was found ──
    row += [_safe(r.get("time_to_best"), 2) for r in exr]
    row += [_safe(get(m, "key_time_to_best"), 2) for m in methods]
    # ── statuses ──
    row += [r.get("status") for r in exr]
    row += [get(m, "key_status") for m in methods]
    # ── pairwise: each exact run vs the first one ──
    for r in exr[1:]:
        row += [_diff(r.get("obj"), exr[0].get("obj")),
                _diff(r.get("lp_relax_ub"), exr[0].get("lp_relax_ub")),
                _ratio(exr[0].get("time"), r.get("time"))]
    # ── pairwise: method vs reference, method vs its own baselines ──
    for m, base in _comparison_pairs(methods):
        row += [_diff(get(m, "key_bound"),  get(base, "key_bound")),
                _diff(get(m, "key_primal"), get(base, "key_primal")),
                _ratio(get(base, "key_time"), get(m, "key_time")),
                _ratio(get(base, "key_time_to_best"), get(m, "key_time_to_best"))]
        if m.get("key_pricing_time") and base.get("key_pricing_time"):
            row.append(_ratio(get(base, "key_pricing_time"), get(m, "key_pricing_time")))
    # ── method-specific extras (grouped by method) ──
    for m in methods:
        for _, rkey in m["extra_cols"]:
            row.append(_safe(results[m["name"]].get(rkey)))
    return row


# ── Instance discovery ────────────────────────────────────────────────────────
_SIZE_PATTERN = re.compile(r"_REQ(\d+)_CDR(\d+)_ODR(\d+)")


def _list_instances(instance_dir, max_req):
    """(nbr_req, nbr_cdr, nbr_odr, name, path) for every .txt file in instance_dir or any of its
    subfolders whose name contains _REQ{n}_CDR{c}_ODR{o} with n ≤ max_req.
    `name` is the path relative to instance_dir (shown in the sheets), `path` the full path.
    Sorted by subfolder, then REQ, CDR and file name, so each subfolder's instances stay together."""
    folder = os.path.join(_REPO_ROOT, instance_dir)
    found = []
    for dirpath, _, filenames in os.walk(folder):
        for filename in filenames:
            match = _SIZE_PATTERN.search(filename)
            if not filename.endswith(".txt") or match is None:
                continue
            nbr_req, nbr_cdr, nbr_odr = map(int, match.groups())
            if nbr_req <= max_req:
                path = os.path.join(dirpath, filename)
                found.append((nbr_req, nbr_cdr, nbr_odr, os.path.relpath(path, folder), path))
    return sorted(found, key=lambda t: (os.path.dirname(t[3]), t[0], t[1], t[3]))


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    dataset  = os.path.basename(os.path.normpath(INSTANCE_DIR))
    cg_tag   = "cgall" if ADD_ALL_COLUMNS else "cgfirst"
    mip_tag  = "_".join(_mode_tag(use_mip) for use_mip in MIP_HEURISTIC_MODES)
    tw_tag   = "_".join(_tw_tag(tighten) for tighten in TIGHTEN_MODES)
    out_xlsx = os.path.join(_HERE, f"bnp_pricing_comparison_{dataset}_{cg_tag}_{mip_tag}_{tw_tag}.xlsx")

    wb       = Workbook()
    hdr_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF")

    def _style(ws):
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font

    ws_exact = None
    first_ws = wb.active
    if RUN_EXACT:
        ws_exact = first_ws
        ws_exact.title = "Exact"
        ws_exact.append(EXACT_HEADER)
        _style(ws_exact)

    method_sheets = {}
    for idx, m in enumerate(METHODS):
        if idx == 0 and not RUN_EXACT:
            ws = first_ws
            ws.title = m["name"]
        else:
            ws = wb.create_sheet(m["name"])
        ws.append(_detail_header(m))
        _style(ws)
        method_sheets[m["name"]] = ws

    comp_header = _comparison_header(METHODS, RUN_EXACT)
    ws_comp = wb.create_sheet("Comparison")
    ws_comp.append(comp_header)
    _style(ws_comp)

    ws_sol = wb.create_sheet("Solution")
    ws_sol.append(SOLUTION_HEADER)
    _style(ws_sol)

    all_sheets = ([ws_exact] if ws_exact else []) + list(method_sheets.values()) + [ws_comp, ws_sol]

    instances = _list_instances(INSTANCE_DIR, MAX_REQ)

    total = len(instances)
    method_names = ", ".join(m["name"] for m in METHODS)
    exact_names  = ", ".join(_exact_names(RUN_EXACT))
    print(f"B&P comparison: {'[' + exact_names + '] + ' if RUN_EXACT else ''}[{method_names}]")
    print(f"  Instances : {total}  ({INSTANCE_DIR}, REQ ≤ {MAX_REQ})")
    print(f"  Output    : {out_xlsx}")
    wb.save(out_xlsx)

    for idx, (nbr_req, nbr_cdr, nbr_odr, instance_file, instance_path) in enumerate(instances, start=1):
        instance_cols = [instance_file, nbr_req, nbr_cdr, nbr_odr]
        print(f"\n{'='*70}")
        print(f"[{idx}/{total}]  REQ={nbr_req}  CDR={nbr_cdr}  ODR={nbr_odr}")
        print(f"  {instance_file}")
        print(f"{'='*70}")

        if not os.path.exists(instance_path):
            print("  [SKIP] file not found")
            if ws_exact:
                for _, tighten in EXACT_RUNS:
                    row = _exact_row(instance_cols, tighten, {})
                    row[EXACT_HEADER.index("Exact_status")] = "FILE NOT FOUND"
                    ws_exact.append(row)
            for m in METHODS:
                hdr = _detail_header(m)
                row = instance_cols + [None] * (len(hdr) - 4)
                row[hdr.index("status")] = "FILE NOT FOUND"
                method_sheets[m["name"]].append(row)
            ws_comp.append(instance_cols + [None] * (len(comp_header) - 4))
            for ws in all_sheets:
                _autosize(ws)
            wb.save(out_xlsx)
            continue

        # ── Exact MIP, once per tightening mode ─────────────────────────────
        res_exact = {}
        if RUN_EXACT:
            for exact_name, tighten in EXACT_RUNS:
                print(f"\n--- {exact_name} ---")
                try:
                    res = solve_exact_mip(instance_path, time_limit=EXACT_TIME_LIMIT,
                                          verbose=False, tighten=tighten)
                except Exception as exc:
                    print(f"  [ERROR] {exc}")
                    res = {"status": f"ERROR: {exc}"}
                res_exact[exact_name] = _flatten_prep(res)
                ws_exact.append(_exact_row(instance_cols, tighten, res))
                print(f"  obj={res.get('obj')}  gap={res.get('mip_gap')}%  "
                      f"time={res.get('time') or 0.0:.1f}s  prep={res.get('prep_time') or 0.0:.2f}s  "
                      f"status={res.get('status')}")
        exact_ref_obj = res_exact.get(EXACT_RUNS[0][0], {}).get("obj") if RUN_EXACT else None

        # ── Configured methods ────────────────────────────────────────────────
        results = {}
        for m in METHODS:
            print(f"\n--- {m['name']} ---")
            try:
                res = m["func"](instance_path, **m["kwargs"])
            except Exception as exc:
                print(f"  [ERROR] {exc}")
                res = {m["key_status"]: f"ERROR: {exc}"}
            results[m["name"]] = _flatten_prep(res)
            method_sheets[m["name"]].append(
                _detail_row(m, res, instance_cols, exact_ref_obj)
            )
            pricing = (f"  pricing={res.get(m['key_pricing_time']) or 0.0:.2f}s"
                       if m.get("key_pricing_time") else "")
            print(f"  primal={res.get(m['key_primal'])}  "
                  f"bound={res.get(m['key_bound']) if m['key_bound'] else None}  "
                  f"time={res.get(m['key_time']) or 0.0:.1f}s{pricing}  "
                  f"status={res.get(m['key_status'])}")

        # ── Comparison row ────────────────────────────────────────────────────
        ws_comp.append(_comparison_row(METHODS, results, res_exact, instance_cols, RUN_EXACT))

        # ── Solution statistics: one row per method (exact MIP first) ─────────
        # Computed on the ORIGINAL instance (not the tightened one), so tw_violations and
        # profit_check also validate the tightening and the original-window costs.
        # Guarded like the solvers: a load or statistics failure becomes an error row
        # instead of stopping the batch before this instance is saved.
        sol_entries = []
        for exact_name in _exact_names(RUN_EXACT):
            res = res_exact.get(exact_name, {})
            sol_entries.append((exact_name, res.get("status"), res.get("obj"), res.get("routes")))
        for m in METHODS:
            res = results[m["name"]]
            sol_entries.append((m["name"], res.get(m["key_status"]), res.get(m["key_primal"]),
                                res.get(m["key_routes"]) if m.get("key_routes") else None))

        try:
            instance, load_error = load_instance_from_txt(instance_path), None
        except Exception as exc:
            instance, load_error = None, exc

        for method_name, status, profit, routes in sol_entries:
            try:
                if load_error is not None:
                    raise load_error
                ws_sol.append(_solution_row(instance, instance_cols, method_name, status,
                                            profit, routes))
            except Exception as exc:
                print(f"  [STATS ERROR] {method_name}: {exc}")
                ws_sol.append(list(instance_cols) + [method_name, f"STATS ERROR: {exc}",
                                                     _safe(profit)] + [None] * len(STAT_KEYS))

        for ws in all_sheets:
            _autosize(ws)
        wb.save(out_xlsx)
        print(f"\n  [SAVED] {out_xlsx}")

    print(f"\n{'='*70}")
    print(f"Done. Results in: {out_xlsx}")
    print(f"{'='*70}")


if __name__ == "__main__":
    main()
