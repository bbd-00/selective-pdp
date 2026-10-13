#!/usr/bin/env python
"""
tw_overlap_visualization.py

For every instance in InstancesTTW/ and InstanceswithDemand/:
  - Saves a Gantt-style time-window plot to  TW_Overlap_Plots/<folder>__<name>.png
  - Writes one summary row to  tw_overlap_statistics.xlsx

Overlap statistics measure how well courier availability windows cover each
request's pickup and delivery time windows.

Run from the project root:
    python tw_overlap_visualization.py
"""

import glob
import os
import re
import sys

import matplotlib
matplotlib.use("Agg")           # no display needed
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def load_instance(filename):
    """
    Tolerant loader that handles both instance families:
      - InstancesTTW:         driver lines have 5 fields (no capacity)
      - InstanceswithDemand:  driver lines have 6 fields (with capacity)
    """
    with open(filename) as f:
        lines = [ln.strip() for ln in f if ln.strip()]
    i = 0

    # use a list as a mutable cursor so nested helpers can advance it
    cur = [0]

    def _next():
        line = lines[cur[0]]; cur[0] += 1
        return line

    assert _next() == "PARAMETERS"
    H, b, W = map(float, _next().split())
    c_com_h, c_com_d, c_occ_o, c_occ_d, rev_order = map(float, _next().split())
    inst = dict(H=H, b=b, W=W, c_com_h=c_com_h, c_com_d=c_com_d,
                c_occ_o=c_occ_o, c_occ_d=c_occ_d, rev_order=rev_order)

    def _parse_drivers(tag):
        assert _next() == tag
        n = int(_next())
        drivers = []
        for _ in range(n):
            p = _next().split()
            did, x, y, e, l = int(p[0]), float(p[1]), float(p[2]), float(p[3]), float(p[4])
            cap = float(p[5]) if len(p) > 5 else None
            drivers.append(dict(id=did, location=(x, y), e=e, l=l, cap=cap))
        return drivers

    inst["committed"]  = _parse_drivers("COMMITTED_DRIVERS")
    inst["occasional"] = _parse_drivers("OCCASIONAL_DRIVERS")

    assert _next() == "REQUESTS"
    n_req = int(_next())
    requests = []
    for _ in range(n_req):
        p = _next().split()
        rid = int(p[0])
        px, py, dx, dy = float(p[1]), float(p[2]), float(p[3]), float(p[4])
        dp, dd = int(p[5]), int(p[6])
        ep, lp, ed, ld = float(p[7]), float(p[8]), float(p[9]), float(p[10])
        dem = float(p[11]) if len(p) > 11 else None
        requests.append(dict(id=rid, pickup=(px, py), dest=(dx, dy),
                             delay_p=dp, delay_d=dd,
                             e_pickup=ep, l_pickup=lp, e_dest=ed, l_dest=ld, l_dem=dem))
    inst["requests"] = requests
    return inst

PLOT_DIR   = "TW_Overlap_Plots"
STATS_FILE = "tw_overlap_statistics.xlsx"

# ── colours ───────────────────────────────────────────────────────────────────
C_CDR   = "#4C72B0"   # committed driver  – blue
C_ODR   = "#DD8452"   # occasional driver – orange
C_PICK  = "#55A868"   # pickup window     – green
C_DELV  = "#C44E52"   # delivery window   – red
BAR_A   = 0.70        # bar alpha

# ── stat columns (order preserved in Excel) ───────────────────────────────────
COLUMNS = [
    "instance", "folder", "n_requests", "n_committed", "n_occasional", "seed", "H",
    # horizon coverage
    "courier_union_coverage",
    "committed_union_coverage",
    "occasional_union_coverage",
    # how many requests have at least one courier that overlaps
    "n_pickups_coverable",
    "n_deliveries_coverable",
    "pct_pickups_coverable",
    "pct_deliveries_coverable",
    # average number of couriers whose window intersects the request TW
    "avg_couriers_per_pickup",
    "avg_couriers_per_delivery",
    # for each request: best (max) overlap / request-TW-length; then averaged
    "avg_best_pickup_overlap_ratio",
    "avg_best_delivery_overlap_ratio",
    "min_best_pickup_overlap_ratio",
    "max_best_pickup_overlap_ratio",
    "min_best_delivery_overlap_ratio",
    "max_best_delivery_overlap_ratio",
]


# ─────────────────────────────────────────────────────────────────────────────
# helpers
# ─────────────────────────────────────────────────────────────────────────────

def _overlap_len(a_e, a_l, b_e, b_l):
    """Length of the intersection of two closed intervals."""
    return max(0.0, min(a_l, b_l) - max(a_e, b_e))


def _union_coverage(windows, H):
    """Fraction of [0, H] covered by the union of (e, l) windows."""
    if not windows:
        return 0.0
    merged = []
    for e, l in sorted((max(0.0, e), min(float(H), l)) for e, l in windows):
        if e >= l:
            continue
        if merged and e <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], l))
        else:
            merged.append((e, l))
    return sum(l - e for e, l in merged) / H


def _per_request_metrics(req_e, req_l, courier_windows):
    """
    Returns (n_couriers_with_overlap, best_overlap_ratio) for one request TW.
    best_overlap_ratio = max_k( |courier_k ∩ req_TW| / |req_TW| )  in [0, 1]
    """
    tw_len = req_l - req_e
    if tw_len <= 0 or not courier_windows:
        return 0, 0.0
    overlaps = [_overlap_len(req_e, req_l, ce, cl) for ce, cl in courier_windows]
    n_overlap  = sum(1 for o in overlaps if o > 0)
    best_ratio = max(overlaps) / tw_len
    return n_overlap, best_ratio


# ─────────────────────────────────────────────────────────────────────────────
# plot
# ─────────────────────────────────────────────────────────────────────────────

def _plot_instance(instance, title, output_path):
    H          = float(instance["H"])
    committed  = instance["committed"]
    occasional = instance["occasional"]
    requests   = instance["requests"]

    # Build rows: (label, e, l, color)  – couriers first, then requests
    rows = []
    for d in committed:
        rows.append((f"CDR {d['id']}", d["e"], d["l"], C_CDR))
    for d in occasional:
        rows.append((f"ODR {d['id']}", d["e"], d["l"], C_ODR))
    n_courier_rows = len(rows)
    for r in requests:
        rows.append((f"R{r['id']} pickup",   r["e_pickup"], r["l_pickup"], C_PICK))
        rows.append((f"R{r['id']} delivery", r["e_dest"],   r["l_dest"],   C_DELV))

    n_rows  = len(rows)
    fig_h   = min(50, max(4, 0.32 * n_rows))
    bar_h   = min(0.7, max(0.15, 0.65 * fig_h / n_rows))
    label_fs = max(4, min(9, int(130 / n_rows)))

    fig, ax = plt.subplots(figsize=(13, fig_h))

    # y positions: row 0 at the top → y = n_rows-1
    for y_rev, (label, e, l, color) in enumerate(rows):
        y = n_rows - 1 - y_rev
        ax.barh(y, l - e, left=e, height=bar_h,
                color=color, alpha=BAR_A, edgecolor="none")

    # dashed separator between couriers and requests
    if 0 < n_courier_rows < n_rows:
        sep_y = n_rows - n_courier_rows - 0.5
        ax.axhline(sep_y, color="grey", linewidth=0.8, linestyle="--", alpha=0.6)

    ax.set_yticks(range(n_rows))
    ax.set_yticklabels([r[0] for r in reversed(rows)], fontsize=label_fs)
    ax.set_xlim(0, H)
    ax.set_xlabel("Time", fontsize=9)
    ax.set_title(title, fontsize=9, pad=6)
    ax.grid(axis="x", linewidth=0.35, alpha=0.45, color="grey")
    ax.tick_params(axis="x", labelsize=8)

    legend_handles = [
        mpatches.Patch(color=C_CDR,  alpha=BAR_A, label="Committed driver"),
        mpatches.Patch(color=C_ODR,  alpha=BAR_A, label="Occasional driver"),
        mpatches.Patch(color=C_PICK, alpha=BAR_A, label="Pickup window"),
        mpatches.Patch(color=C_DELV, alpha=BAR_A, label="Delivery window"),
    ]
    ax.legend(handles=legend_handles, loc="upper right", fontsize=7.5, ncol=2,
              framealpha=0.85)

    plt.tight_layout()
    plt.savefig(output_path, dpi=130, bbox_inches="tight")
    plt.close(fig)


# ─────────────────────────────────────────────────────────────────────────────
# statistics
# ─────────────────────────────────────────────────────────────────────────────

def _compute_stats(instance, instance_name, folder):
    H          = float(instance["H"])
    committed  = instance["committed"]
    occasional = instance["occasional"]
    requests   = instance["requests"]
    all_couriers = committed + occasional
    courier_wins = [(d["e"], d["l"]) for d in all_couriers]

    pickup_n, pickup_r, delivery_n, delivery_r = [], [], [], []
    for req in requests:
        cn, cr = _per_request_metrics(req["e_pickup"], req["l_pickup"], courier_wins)
        pickup_n.append(cn)
        pickup_r.append(cr)

        dn, dr = _per_request_metrics(req["e_dest"], req["l_dest"], courier_wins)
        delivery_n.append(dn)
        delivery_r.append(dr)

    n_req = len(requests)
    safe_mean = lambda lst: float(np.mean(lst)) if lst else 0.0
    safe_min  = lambda lst: float(np.min(lst))  if lst else 0.0
    safe_max  = lambda lst: float(np.max(lst))  if lst else 0.0

    m = re.search(r"REQ(\d+)_CDR(\d+)_ODR(\d+)_FS\d+_SS(\d+)", instance_name)
    n_req_p = int(m.group(1)) if m else n_req
    n_cdr_p = int(m.group(2)) if m else len(committed)
    n_odr_p = int(m.group(3)) if m else len(occasional)
    seed    = int(m.group(4)) if m else -1

    n_pick_cov = sum(1 for n in pickup_n if n > 0)
    n_delv_cov = sum(1 for n in delivery_n if n > 0)

    return {
        "instance":                        f"{folder}/{instance_name}",
        "folder":                          folder,
        "n_requests":                      n_req_p,
        "n_committed":                     n_cdr_p,
        "n_occasional":                    n_odr_p,
        "seed":                            seed,
        "H":                               int(H),
        "courier_union_coverage":          round(_union_coverage(courier_wins, H), 4),
        "committed_union_coverage":        round(_union_coverage([(d["e"], d["l"]) for d in committed], H), 4),
        "occasional_union_coverage":       round(_union_coverage([(d["e"], d["l"]) for d in occasional], H), 4),
        "n_pickups_coverable":             n_pick_cov,
        "n_deliveries_coverable":          n_delv_cov,
        "pct_pickups_coverable":           round(100 * n_pick_cov / n_req, 2) if n_req else 0.0,
        "pct_deliveries_coverable":        round(100 * n_delv_cov / n_req, 2) if n_req else 0.0,
        "avg_couriers_per_pickup":         round(safe_mean(pickup_n), 3),
        "avg_couriers_per_delivery":       round(safe_mean(delivery_n), 3),
        "avg_best_pickup_overlap_ratio":   round(safe_mean(pickup_r), 4),
        "avg_best_delivery_overlap_ratio": round(safe_mean(delivery_r), 4),
        "min_best_pickup_overlap_ratio":   round(safe_min(pickup_r), 4),
        "max_best_pickup_overlap_ratio":   round(safe_max(pickup_r), 4),
        "min_best_delivery_overlap_ratio": round(safe_min(delivery_r), 4),
        "max_best_delivery_overlap_ratio": round(safe_max(delivery_r), 4),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Excel writer
# ─────────────────────────────────────────────────────────────────────────────

def _write_excel(rows, path):
    wb  = Workbook()
    ws  = wb.active
    ws.title = "TW Overlap Statistics"

    hdr_fill = PatternFill("solid", fgColor="2F5496")
    hdr_font = Font(color="FFFFFF", bold=True)

    for col_i, name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_i, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_i)].width = max(len(name) + 2, 12)

    for row_i, row in enumerate(rows, 2):
        for col_i, name in enumerate(COLUMNS, 1):
            ws.cell(row=row_i, column=col_i, value=row.get(name, ""))

    ws.freeze_panes = "A2"
    wb.save(path)


# ─────────────────────────────────────────────────────────────────────────────
# main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    root     = os.path.dirname(os.path.abspath(__file__))
    plot_dir = os.path.join(root, PLOT_DIR)
    os.makedirs(plot_dir, exist_ok=True)

    instance_dirs = ["InstancesTTW", "InstanceswithDemand"]
    all_paths = []
    for d in instance_dirs:
        all_paths.extend(sorted(glob.glob(os.path.join(root, d, "*.txt"))))

    print(f"Found {len(all_paths)} instances across {instance_dirs}.")

    stats_rows = []
    for path in all_paths:
        folder   = os.path.basename(os.path.dirname(path))
        basename = os.path.splitext(os.path.basename(path))[0]
        label    = f"{folder}/{basename}"
        print(f"  {label} ...", end=" ", flush=True)

        try:
            instance = load_instance(path)
        except Exception as exc:
            print(f"SKIP ({exc})")
            continue

        safe = f"{folder}__{basename}"
        _plot_instance(instance, label, os.path.join(plot_dir, f"{safe}.png"))

        stats_rows.append(_compute_stats(instance, basename, folder))
        print("done")

    stats_path = os.path.join(root, STATS_FILE)
    _write_excel(stats_rows, stats_path)

    print(f"\nPlots  → {plot_dir}/  ({len(stats_rows)} files)")
    print(f"Stats  → {stats_path}")


if __name__ == "__main__":
    main()
